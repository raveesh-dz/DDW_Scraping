"""
Compare the two most recent DDW scraped Excel files and produce a highlighted
output that marks newly added sessions and sub-sessions in green.

Usage::

    python compare_excel.py                  # auto-detects the two latest files
    python compare_excel.py OLD.xlsx NEW.xlsx  # explicit paths
"""

import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

RESULTS_DIR = Path(__file__).resolve().parent / "results"
HEADER_ROW = 2          # headers live in row 2
DATA_START  = HEADER_ROW + 1

NEW_FILL   = PatternFill(fill_type="solid", fgColor="C6EFCE")   # light green
NEW_FILL_S = PatternFill(fill_type="solid", fgColor="92D050")   # darker green for session rows


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def latest_result_files():
    """Return (older, newer) non-checkpoint xlsx paths sorted by name."""
    files = sorted(
        p for p in RESULTS_DIR.glob("ddw-conference-schedule-2026-*.xlsx")
        if "_checkpoint" not in p.name
        and not p.name.startswith("~$")
    )
    if len(files) < 2:
        raise FileNotFoundError(
            f"Need at least 2 result files in {RESULTS_DIR}, found {len(files)}"
        )
    return files[-2], files[-1]


def get_headers(ws):
    """Return column-name → 1-based column-index mapping from HEADER_ROW."""
    return {
        cell.value: cell.column
        for cell in ws[HEADER_ROW]
        if cell.value is not None
    }


def row_to_dict(row_values, headers):
    """Map a tuple of row values to a {header: value} dict."""
    result = {}
    for name, col_idx in headers.items():
        result[name] = row_values[col_idx - 1] if col_idx - 1 < len(row_values) else None
    return result


def session_key(d):
    """Unique key for a session (non-None Session Name row)."""
    return (
        str(d.get("Session Name") or "").strip(),
        str(d.get("Session Start Time") or "").strip(),
        str(d.get("Location") or "").strip(),
    )


def subsession_key(d, parent_session_name):
    """Unique key for a sub-session row."""
    return (
        parent_session_name,
        str(d.get("Sub-Session Title") or "").strip(),
        str(d.get("Sub-Session Start") or "").strip(),
    )


def extract_keys(ws):
    """
    Walk a worksheet and return:
      - session_keys: set of session_key tuples
      - subsession_keys: set of subsession_key tuples
    """
    headers = get_headers(ws)
    session_col  = headers.get("Session Name")
    sub_col      = headers.get("Sub-Session Title")

    session_keys    = set()
    subsession_keys = set()
    current_session = ""

    for row_idx in range(DATA_START, ws.max_row + 1):
        row = tuple(cell.value for cell in ws[row_idx])
        d   = row_to_dict(row, headers)

        sname = str(d.get("Session Name") or "").strip()
        if sname:
            current_session = sname
            key = session_key(d)
            if any(k for k in key):
                session_keys.add(key)

        sub = str(d.get("Sub-Session Title") or "").strip()
        if sub:
            key = subsession_key(d, current_session)
            if any(k for k in key):
                subsession_keys.add(key)

    return session_keys, subsession_keys


# ---------------------------------------------------------------------------
# Main comparison
# ---------------------------------------------------------------------------

def compare_and_highlight(old_path: Path, new_path: Path, out_path: Path):
    print(f"Older : {old_path.name}")
    print(f"Newer : {new_path.name}")

    wb_new = load_workbook(new_path)

    total_new_sessions    = 0
    total_new_subsessions = 0

    for sheet_name in wb_new.sheetnames:
        ws_new = wb_new[sheet_name]

        # Try to open matching sheet in old file
        wb_old = load_workbook(old_path)
        if sheet_name in wb_old.sheetnames:
            ws_old = wb_old[sheet_name]
            old_session_keys, old_subsession_keys = extract_keys(ws_old)
        else:
            # Sheet didn't exist before → everything is new
            old_session_keys    = set()
            old_subsession_keys = set()

        headers = get_headers(ws_new)
        max_col = ws_new.max_column

        new_sessions_sheet    = 0
        new_subsessions_sheet = 0
        current_session       = ""
        current_session_new   = False

        for row_idx in range(DATA_START, ws_new.max_row + 1):
            row = tuple(cell.value for cell in ws_new[row_idx])
            d   = row_to_dict(row, headers)

            sname = str(d.get("Session Name") or "").strip()
            sub   = str(d.get("Sub-Session Title") or "").strip()

            is_new_session    = False
            is_new_subsession = False

            if sname:
                current_session = sname
                key = session_key(d)
                if any(k for k in key) and key not in old_session_keys:
                    is_new_session     = True
                    current_session_new = True
                    new_sessions_sheet += 1
                else:
                    current_session_new = False

            if sub:
                key = subsession_key(d, current_session)
                if any(k for k in key) and key not in old_subsession_keys:
                    is_new_subsession      = True
                    new_subsessions_sheet += 1

            fill = None
            if is_new_session:
                fill = NEW_FILL_S
            elif is_new_subsession or current_session_new:
                fill = NEW_FILL

            if fill:
                for col in range(1, max_col + 1):
                    ws_new.cell(row=row_idx, column=col).fill = fill

        total_new_sessions    += new_sessions_sheet
        total_new_subsessions += new_subsessions_sheet
        print(
            f"  {sheet_name}: {new_sessions_sheet} new sessions, "
            f"{new_subsessions_sheet} new sub-sessions highlighted"
        )

    wb_new.save(out_path)
    print(f"\nSaved: {out_path}")
    print(
        f"Total: {total_new_sessions} new sessions, "
        f"{total_new_subsessions} new sub-sessions"
    )


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) == 3:
        old_path = Path(sys.argv[1])
        new_path = Path(sys.argv[2])
    else:
        old_path, new_path = latest_result_files()

    out_path = new_path.with_name(new_path.stem + "_diff.xlsx")
    compare_and_highlight(old_path, new_path, out_path)


if __name__ == "__main__":
    main()
