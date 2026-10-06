"""
retag_from_excel.py — Re-apply tagging logic to an existing scraped Excel file.

Reads an already-scraped DDW Excel (output of scrape_ddw_schedule.py), rebuilds
the session/sub-session structure from it, runs the updated keyword-matching logic
(per-sub-session tagging, Unicode normalisation, substring deduplication), and
writes a new Excel with corrected tags.

This lets you QC the tagging fixes without re-running the 3-hour browser scrape.

Usage::

    python retag_from_excel.py                       # uses latest result file
    python retag_from_excel.py path/to/input.xlsx    # explicit input
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ddw_scraper import config
from ddw_scraper import keywords as kw
from ddw_scraper.exporter import write_excel
from ddw_scraper.logger import log

# ---------------------------------------------------------------------------
# Column positions (1-based, matching exporter._COLUMNS layout)
# ---------------------------------------------------------------------------
COL_INDICATION   = 3   # C
COL_COMPANY      = 4   # D
COL_PRODUCT      = 5   # E
COL_MOA          = 6   # F
COL_TRACK        = 7   # G
COL_SCIENCE_TYPE = 8   # H
COL_SESSION_NAME = 9   # I  (merged — None on continuation rows)
COL_FORMAT       = 10  # J  (merged)
COL_START        = 11  # K  (merged)
COL_END          = 12  # L  (merged)
COL_LOCATION     = 13  # M  (merged)
COL_SUB_START    = 15  # O
COL_SUB_END      = 16  # P
COL_SUB_NAME     = 17  # Q
COL_SOCIETY      = 19  # S


def _latest_result() -> Path:
    files = sorted(
        p for p in config.RESULTS_DIR.glob("ddw-conference-schedule-2026-*.xlsx")
        if "_checkpoint" not in p.name
        and "_retagged" not in p.name
        and not p.name.startswith("~$")
    )
    if not files:
        raise FileNotFoundError(f"No result files found in {config.RESULTS_DIR}")
    return files[-1]


def _read_sessions(wb) -> list[dict]:
    """Parse an existing Excel workbook into session row dicts."""
    import openpyxl

    rows = []

    for sheet_name, day_label in config.DAY_SHEETS:
        if sheet_name not in wb.sheetnames:
            log(f"Sheet '{sheet_name}' not found, skipping.")
            continue

        ws = wb[sheet_name]

        # Carry-forward state for merged session-level cells
        cur_session  = ""
        cur_format   = ""
        cur_start    = ""
        cur_end      = ""
        cur_location = ""
        cur_row: dict | None = None

        for row_idx in range(3, ws.max_row + 1):
            def v(col):
                val = ws.cell(row_idx, col).value
                return str(val).strip() if val is not None else ""

            session_name = v(COL_SESSION_NAME)
            sub_name     = v(COL_SUB_NAME)

            if not session_name and not sub_name:
                continue  # blank spacer row

            # Session-level fields — update carry-forward when a new session starts
            if session_name:
                cur_session  = session_name
                cur_format   = v(COL_FORMAT)
                cur_start    = v(COL_START)
                cur_end      = v(COL_END)
                cur_location = v(COL_LOCATION)
                cur_row = {
                    "day_label":    day_label,
                    "session_name": cur_session,
                    "start_time":   cur_start,
                    "end_time":     cur_end,
                    "date":         day_label,
                    "region":       cur_location,
                    "track":        v(COL_TRACK),
                    "format":       cur_format,
                    "society":      v(COL_SOCIETY),
                    "science_type": v(COL_SCIENCE_TYPE),
                    "notes":        "",
                    "sub_sessions": [],
                    "_session_kw":  {},
                }
                rows.append(cur_row)

            # Track/Science Type/Society repeat every row (not merged)
            # Update the current session row so we always have the latest values
            if cur_row is not None:
                track = v(COL_TRACK)
                sci   = v(COL_SCIENCE_TYPE)
                soc   = v(COL_SOCIETY)
                if track: cur_row["track"]        = track
                if sci:   cur_row["science_type"] = sci
                if soc:   cur_row["society"]      = soc

            # Sub-session row
            if sub_name and cur_row is not None:
                cur_row["sub_sessions"].append({
                    "name":  sub_name,
                    "start": v(COL_SUB_START),
                    "end":   v(COL_SUB_END),
                })

    return rows


def _apply_tags(rows: list[dict], patterns: dict) -> None:
    """Run per-sub-session keyword matching in-place on all rows."""
    for r in rows:
        session_base = [
            r["session_name"],
            r["track"],
            r["format"],
            r["society"],
            r["science_type"],
        ]

        if r["sub_sessions"]:
            for sub in r["sub_sessions"]:
                sub_kw = kw.match_keywords(session_base + [sub["name"]], patterns)
                sub["indication"]   = sub_kw.get("Indication", "")
                sub["company"]      = sub_kw.get("Company", "")
                sub["product_name"] = sub_kw.get("Product Name", "")
                sub["moa"]          = sub_kw.get("MOA", "")
        else:
            # No sub-sessions: compute tags from session-level fields only
            session_kw = kw.match_keywords(session_base, patterns)
            r["_session_kw"] = session_kw


def main() -> int:
    import openpyxl

    if len(sys.argv) >= 2:
        input_path = Path(sys.argv[1])
    else:
        input_path = _latest_result()

    output_path = input_path.with_name(input_path.stem + "_retagged.xlsx")

    log(f"Input : {input_path.name}")
    log(f"Output: {output_path.name}")

    log("Loading keywords...")
    patterns = kw.build_patterns(kw.load_keywords())
    total_kw = sum(len(v) for v in patterns.values())
    log(f"  {total_kw} keywords across {len(patterns)} columns.")

    log("Reading scraped Excel...")
    wb = openpyxl.load_workbook(input_path, data_only=True)
    rows = _read_sessions(wb)
    wb.close()
    log(f"  {len(rows)} sessions read.")

    log("Applying per-sub-session tags...")
    _apply_tags(rows, patterns)

    total_subs = sum(len(r["sub_sessions"]) for r in rows)
    tagged_subs = sum(
        1 for r in rows for sub in r["sub_sessions"]
        if any([sub.get("indication"), sub.get("company"),
                sub.get("product_name"), sub.get("moa")])
    )
    log(f"  {tagged_subs} / {total_subs} sub-session rows have at least one tag.")

    log("Writing output Excel...")
    write_excel(rows, output_path)
    log("Done.")
    return 0


if __name__ == "__main__":
    exit(main())
