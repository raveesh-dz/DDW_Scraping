import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

f = Path(__file__).parent / "ats_output" / "ATS_Abstracts 2026 2.xlsx"
wb = load_workbook(f, read_only=True)
print("SHEETS:", wb.sheetnames)
for s in wb.sheetnames:
    df = pd.read_excel(f, sheet_name=s, nrows=5)
    print(f"\n=== {s} ===")
    print("COLS:", list(df.columns))
    print(df.to_string())
    full = pd.read_excel(f, sheet_name=s)
    print(f"ROWS: {len(full)}")
